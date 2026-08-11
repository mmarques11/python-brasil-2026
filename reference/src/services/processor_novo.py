import asyncio
from pathlib import Path
from typing import TypeAlias

from openpyxl import load_workbook
from src.plugins.driver import DriverPlaywright


WORKER_COUNT = 3
ROUND_MAX = 10
RPA_CHALLENGE_URL = 'https://rpachallenge.com/'
DATA_FILE = Path(__file__).resolve().parents[2] / 'data' / 'dados.xlsx'

DadosFormulario: TypeAlias = dict[str, str]

COLUNAS_OBRIGATORIAS = (
    'First Name',
    'Last Name',
    'Company Name',
    'Role in Company',
    'Address',
    'Email',
    'Phone Number',
)


def _converter_para_texto(valor: object) -> str:
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def carregar_dados() -> list[DadosFormulario]:
    if not DATA_FILE.exists():
        raise FileNotFoundError(f'Planilha não encontrada: {DATA_FILE}')

    workbook = load_workbook(DATA_FILE, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = [_converter_para_texto(value) for value in next(rows)]

        colunas_faltantes = set(COLUNAS_OBRIGATORIAS).difference(headers)
        if colunas_faltantes:
            colunas = ', '.join(sorted(colunas_faltantes))
            raise ValueError(f'Colunas obrigatórias ausentes: {colunas}')

        indice_por_coluna = {
            nome_coluna: headers.index(nome_coluna)
            for nome_coluna in COLUNAS_OBRIGATORIAS
        }
        dados = [
            {
                nome_coluna: _converter_para_texto(row[indice])
                for nome_coluna, indice in indice_por_coluna.items()
            }
            for row in rows
            if any(value is not None for value in row)
        ]
    finally:
        workbook.close()

    if len(dados) != ROUND_MAX:
        raise ValueError(
            f'A planilha deve possuir {ROUND_MAX} registros; '
            f'foram encontrados {len(dados)}.'
        )

    return dados


def dividir_dados(
    dados: list[DadosFormulario],
    quantidade_workers: int,
) -> list[list[DadosFormulario]]:
    if quantidade_workers <= 0:
        raise ValueError('A quantidade de workers deve ser maior que zero.')

    tamanho_base = len(dados) // quantidade_workers
    lotes: list[list[DadosFormulario]] = []

    for indice in range(quantidade_workers):
        inicio = indice * tamanho_base
        fim = (
            len(dados)
            if indice == quantidade_workers - 1
            else inicio + tamanho_base
        )
        lotes.append(dados[inicio:fim])

    return lotes


async def executar_worker(
    driver: DriverPlaywright,
    worker_id: int,
    dados: list[DadosFormulario],
) -> None:
    context = await driver.novo_contexto()

    try:
        page = await driver.nova_pagina(context)
        await page.goto(RPA_CHALLENGE_URL)
        await page.get_by_role('button', name='Start').click()

        for rodada, dados_rodada in enumerate(dados, start=1):
            for nome_campo in COLUNAS_OBRIGATORIAS:
                campo = page.locator(
                    'label',
                    has_text=nome_campo,
                ).locator('xpath=following-sibling::input')
                await campo.fill(dados_rodada[nome_campo])

            await page.get_by_role('button', name='Submit').click()
            print(f'Worker {worker_id}: item {rodada}/{len(dados)} do seu lote.')

        print(f'Worker {worker_id}: processamento concluído.')
    finally:
        await context.close()


async def executar_processamento() -> None:
    dados = carregar_dados()
    lotes = dividir_dados(dados, WORKER_COUNT)
    driver = DriverPlaywright(headless=False)
    await driver.iniciar_driver()

    try:
        await asyncio.gather(
            *(
                executar_worker(driver, worker_id, lote)
                for worker_id, lote in enumerate(lotes, start=1)
            )
        )
    finally:
        await driver.fechar_driver()
